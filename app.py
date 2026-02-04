import streamlit as st
import requests
import json
import sqlite3
import msgpack
import time
import pandas as pd
import hashlib
import docx  # NEW: For Word Docs
from PyPDF2 import PdfReader
from langchain_community.embeddings import OllamaEmbeddings
from langchain_community.llms import Ollama

# --- FIX: SMART IMPORT FOR TEXT SPLITTER ---
try:
    from langchain_text_splitters import RecursiveCharacterTextSplitter
except ImportError:
    from langchain.text_splitter import RecursiveCharacterTextSplitter

# --- SERVER CONFIG ---
ENDEE_URL = "http://127.0.0.1:8080/api/v1"
AUTH_TOKEN = "secret123"
SQL_DB = "rag_memory.db"
BATCH_SIZE = 100 

# --- MODELS (LOCAL & FAST) ---
embeddings = OllamaEmbeddings(model="nomic-embed-text")
llm = Ollama(model="llama3.2:1b", keep_alive="1h")

# --- DATABASE LAYER ---
def init_text_store():
    with sqlite3.connect(SQL_DB) as conn:
        conn.execute('''
            CREATE TABLE IF NOT EXISTS user_docs (
                username TEXT,
                id TEXT, 
                content TEXT,
                PRIMARY KEY (username, id)
            )
        ''')
        conn.execute("PRAGMA journal_mode = WAL;")
        conn.execute("PRAGMA synchronous = NORMAL;")

def save_text_batch(username, data_list):
    with sqlite3.connect(SQL_DB) as conn:
        secure_data = [(username, d[0], d[1]) for d in data_list]
        conn.executemany("INSERT OR REPLACE INTO user_docs VALUES (?, ?, ?)", secure_data)

def get_text_map(username, id_list):
    with sqlite3.connect(SQL_DB) as conn:
        if not id_list: return {}
        placeholders = ',' .join('?' for _ in id_list)
        query = f"SELECT id, content FROM user_docs WHERE username = ? AND id IN ({placeholders})"
        args = [username] + [str(x) for x in id_list]
        rows = conn.execute(query, args).fetchall()
        return {r[0]: r[1] for r in rows}

def flush_user_data(username):
    with sqlite3.connect(SQL_DB) as conn:
        conn.execute("DELETE FROM user_docs WHERE username = ?", (username,))
        conn.commit()

class EndeeClient:
    def __init__(self, base_url, token):
        self.url = base_url
        self.headers = {"Authorization": token, "Content-Type": "application/json"}

    def _post(self, endpoint, payload):
        try:
            r = requests.post(f"{self.url}{endpoint}", json=payload, headers=self.headers, timeout=30)
            if r.status_code == 200 and not r.content: return {"success": True}
            if r.headers.get("Content-Type") == "application/msgpack":
                return msgpack.unpackb(r.content, raw=False)
            return r.json()
        except Exception as e:
            return {"error": str(e)}

    def create_index(self, name):
        return self._post("/index/create", {
            "index_name": name, "dim": 768, "space_type": "cosine", "M": 32, "ef_con": 200
        })

    def delete_index(self, name):
        try:
            requests.delete(f"{self.url}/index/{name}/delete", headers=self.headers, timeout=5)
            return True
        except: return False

    def insert(self, idx_name, vectors):
        return self._post(f"/index/{idx_name}/vector/insert", vectors)

    def search(self, idx_name, vec, k=5):
        return self._post(f"/index/{idx_name}/search", {"vector": vec, "k": k, "ef": 50})

    def info(self, idx_name):
        try:
            r = requests.get(f"{self.url}/index/{idx_name}/info", headers=self.headers)
            return r.json()
        except: return {}

db = EndeeClient(ENDEE_URL, AUTH_TOKEN)

# --- PROCESSING ENGINE (OMNI-FORMAT) ---
def read_file_content(uploaded_file):
    """Smart text extraction based on file extension"""
    name = uploaded_file.name.lower()
    text = ""
    
    try:
        # 1. PDF
        if name.endswith(".pdf"):
            reader = PdfReader(uploaded_file)
            text = "".join([p.extract_text() for p in reader.pages if p.extract_text()])
            
        # 2. WORD (.docx)
        elif name.endswith(".docx"):
            doc = docx.Document(uploaded_file)
            text = "\n".join([p.text for p in doc.paragraphs])
            
        # 3. TEXT / MARKDOWN (.txt, .md)
        elif name.endswith(".txt") or name.endswith(".md"):
            # We must decode bytes to string
            text = uploaded_file.read().decode("utf-8")
            
    except Exception as e:
        st.error(f"Error reading {name}: {e}")
        return ""
        
    return text

def process_uploaded_files(uploaded_files, username, idx_name):
    text_splitter = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=100)
    
    total_docs = []
    status_text = st.empty()
    progress_bar = st.progress(0)
    
    status_text.text("📖 Reading files...")
    
    # Reading Phase
    for f in uploaded_files:
        text = read_file_content(f)
        if text:
            chunks = text_splitter.split_text(text)
            for i, chunk in enumerate(chunks):
                safe_id = f"{username}_{f.name}_{i}"
                total_docs.append({"id": safe_id, "text": chunk})
    
    if not total_docs:
        status_text.warning("No text extracted. Are files empty?")
        return

    # Batch Processing Phase
    total_chunks = len(total_docs)
    status_text.text(f"🚀 Vectorizing {total_chunks} chunks...")
    
    for i in range(0, total_chunks, BATCH_SIZE):
        batch = total_docs[i : i + BATCH_SIZE]
        texts_to_embed = [d["text"] for d in batch]
        vectors = embeddings.embed_documents(texts_to_embed)
        
        vec_batch = []
        txt_batch = []
        
        for j, vec in enumerate(vectors):
            doc = batch[j]
            vec_batch.append({"id": doc["id"], "vector": vec})
            txt_batch.append((doc["id"], doc["text"]))
        
        save_text_batch(username, txt_batch)
        db.insert(idx_name, vec_batch)
        
        progress = min((i + BATCH_SIZE) / total_chunks, 1.0)
        progress_bar.progress(progress)

    status_text.text("✅ All files processed successfully!")
    time.sleep(1)
    status_text.empty()
    progress_bar.empty()

# --- UI LOGIC ---
st.set_page_config(page_title="Endee Turbo RAG", layout="wide")
st.title("🚀 Endee RAG: All Formats")

init_text_store()

if 'username' not in st.session_state:
    st.session_state.username = None

if not st.session_state.username:
    col1, col2, col3 = st.columns([1,2,1])
    with col2:
        st.subheader("Login")
        user_input = st.text_input("Username")
        if st.button("Start Session"):
            if user_input:
                st.session_state.username = user_input.strip()
                st.rerun()
else:
    user_hash = hashlib.md5(st.session_state.username.encode()).hexdigest()[:8]
    idx_name = f"idx_{user_hash}_v1"
    db.create_index(idx_name)
    
    with st.sidebar:
        st.success(f"User: **{st.session_state.username}**")
        st.info("⚡ Mode: Omni-Format")
        
        st.divider()
        k_val = st.slider("Context Chunks", 1, 15, 5)
        min_score = st.slider("Min Relevance", 0.0, 1.0, 0.35)
        
        stats = db.info(idx_name)
        count = stats.get('total_elements', 0) 
        st.metric("Total Vectors", count)
        
        st.divider()
        if st.button("Flush Data (Reset)", type="primary"):
            db.delete_index(idx_name)
            flush_user_data(st.session_state.username)
            st.rerun()

        if st.button("Logout"):
            st.session_state.username = None
            st.rerun()

    tab1, tab2 = st.tabs(["📂 Multi-File Upload", "⚡ Fast Search"])
    
    with tab1:
        # UPDATED: File uploader accepts more types
        files = st.file_uploader(
            "Upload Documents (PDF, Word, Text, Markdown)", 
            type=["pdf", "docx", "txt", "md"], 
            accept_multiple_files=True
        )
        if files and st.button(f"Process {len(files)} Files"):
            process_uploaded_files(files, st.session_state.username, idx_name)
            st.rerun()

    with tab2:
        q = st.text_input("Ask your documents")
        if q:
            q_vec = embeddings.embed_query(q)
            results = db.search(idx_name, q_vec, k=k_val)
            
            parsed_data = []
            if isinstance(results, list):
                for item in results:
                    if isinstance(item, list) and len(item) > 1:
                        score = float(item[0])
                        raw_id = item[1]
                        if score >= min_score:
                            doc_id = raw_id.decode('utf-8') if isinstance(raw_id, bytes) else str(raw_id)
                            parsed_data.append({"ID": doc_id, "Score": score})
            elif isinstance(results, dict):
                items = results.get("matches") or results.get("results") or []
                for x in items:
                    score = float(x.get("score", 0))
                    if score >= min_score:
                        parsed_data.append({"ID": x.get("id"), "Score": score})

            if parsed_data:
                parsed_data.sort(key=lambda x: x["Score"], reverse=True)
                text_map = get_text_map(st.session_state.username, [row["ID"] for row in parsed_data])
                
                context_str = ""
                for i, row in enumerate(parsed_data):
                    full_text = text_map.get(row["ID"], "")
                    context_str += f"--- SOURCE {i+1} ---\n{full_text}\n\n"

                with st.expander(f"See {len(parsed_data)} Matches used"):
                     st.text(context_str[:2000] + "...")

                prompt = f"""
You are a fast analysis engine. 
Based STRICTLY on the context below, answer the user's question.
If the answer is not in the text, simply say "Not found in documents."
Do not refuse. Do not lecture. Be concise.

CONTEXT:
{context_str}

QUESTION: 
{q}

ANSWER:
"""
                st.divider()
                ans_box = st.empty()
                full_ans = ""
                
                try:
                    for chunk in llm.stream(prompt):
                        full_ans += chunk
                        ans_box.markdown(f"**Answer:** {full_ans}▌")
                    ans_box.markdown(f"**Answer:** {full_ans}")
                except Exception as e:
                    st.error(f"Error: {e}")
            else:
                st.warning("No relevant info found.")
